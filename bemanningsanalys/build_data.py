#!/usr/bin/env python3
"""Bygger data/preloaded.js från Trafikverkets underlag (uppdrag bemanning 2027, VO PR).

Användning:  python3 -I build_data.py <katalog med xlsx/pptx-filerna>
Filerna identifieras på delar av filnamnet, så prefix från uppladdning spelar ingen roll.
"""
import glob
import json
import os
import re
import sys
import warnings
from datetime import datetime

import openpyxl

warnings.filterwarnings("ignore")

SRC = sys.argv[1] if len(sys.argv) > 1 else "."
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "preloaded.js")

AVD = {
    "PR": "Stora projekt (PR)", "PR (L)": "Stora projekt Ledning", "PRh": "HR och Säkerhet",
    "PRve": "Verksamhetsstyrning", "PRn": "Järnväg Nord", "PRs": "Järnväg Syd",
    "PRo": "Ostlänken", "PRv": "Väg och vatten", "PRö": "Väg/Järnväg Öst", "PRvl": "Västlänken",
}
KAT_FIX = {  # trunkerade namn i åtgärdsplanens validering -> fullständigt
    "Teknik, miljö och markförhandl": "Teknik, miljö och markförhandling",
}

def find(part):
    m = [f for f in glob.glob(os.path.join(SRC, "*")) if part.lower() in os.path.basename(f).lower()]
    if not m:
        raise SystemExit(f"Hittar ingen fil med '{part}' i {SRC}")
    return m[0]

def num(x):
    if isinstance(x, (int, float)) and not isinstance(x, bool):
        return round(float(x), 4)
    return None

def s(x):
    return None if x is None else str(x).strip() or None

def iso(x):
    if isinstance(x, datetime):
        return x.strftime("%Y-%m")
    t = s(x)
    if not t:
        return None
    sv = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "maj": 5, "juni": 6, "jun": 6, "juli": 7, "jul": 7,
          "aug": 8, "sep": 9, "okt": 10, "nov": 11, "dec": 12}
    m = re.match(r"([a-zåäö]+)-(\d{2})$", t.lower())
    if m and m.group(1) in sv:
        return f"20{m.group(2)}-{sv[m.group(1)]:02d}"
    if re.fullmatch(r"\d{4}", t):
        return t
    return t

DATASETS = []

def add(id_, name, source, desc, cols, rows):
    DATASETS.append({"id": id_, "name": name, "source": source, "desc": desc,
                     "columns": cols, "rows": rows})
    print(f"  {id_}: {len(rows)} rader")

def col(name, kind="text", role=None):
    return {"name": name, "type": kind, **({"role": role} if role else {})}

# ---------------------------------------------------------------- åtgärdsplan
f_plan = find("tg_rdsplan")
wb = openpyxl.load_workbook(f_plan, data_only=True)
ws = wb["Åtgärdsplan"]
src = os.path.basename(f_plan)
rows = []
for r in range(27, ws.max_row + 1):
    avd = s(ws.cell(r, 2).value)
    if not avd:
        continue
    kat = s(ws.cell(r, 4).value)
    kat = KAT_FIX.get(kat, kat)
    egna, pr, trv = (num(ws.cell(r, c).value) or 0 for c in (7, 8, 9))
    rows.append([avd, AVD.get(avd, avd), s(ws.cell(r, 3).value), kat, s(ws.cell(r, 5).value),
                 iso(ws.cell(r, 6).value), egna, pr, trv, round(egna + pr + trv, 4)])
add("atgarder", "Åtgärder (uppdrag bemanning 2027)", src,
    "Föreslagna åtgärder per avdelning med effekt i ÅA (negativt tal = minskning). Använd för frågor om hur många ÅA vi minskar med.",
    [col("Avdelningskod", role="dim"), col("Avdelning", role="dim"), col("Åtgärd", role="dim"),
     col("Befattningskategori", role="dim"), col("Konsekvens", role="text"), col("Tidplan", role="dim"),
     col("Egna resurser ÅA", "num"), col("Inlånade internt PR ÅA", "num"),
     col("Inlånade övriga TRV ÅA", "num"), col("Total förändring ÅA", "num")], rows)

# sammanfattning per kategori (PR totalt)
rows = []
hdr = [ws.cell(9, c).value for c in range(4, 22)]
for r in range(10, 21):
    k = s(ws.cell(r, 3).value)
    if not k:
        continue
    k = KAT_FIX.get(k, k)
    rows.append([k] + [num(ws.cell(r, c).value) or 0 for c in (4, 5, 6, 7, 8, 9, 11, 12, 13, 15, 16, 17, 19, 20, 21)])
add("atgardsplan_summering", "Åtgärdsplan – summering per befattningskategori (VO PR)", src,
    "Basnivå T1, resursbehov T2, effekt av åtgärder samt Heroma-utfall (MA) och inlånade resurser per kategori.",
    [col("Befattningskategori", role="dim")] + [col(n, "num") for n in [
        "Basnivå T1 2026 (ÅA)", "Resursbehov T2 2026 (ÅA)", "Effekt egna resurser (ÅA)",
        "Effekt inlånade internt PR (ÅA)", "Effekt inlånade övriga TRV (ÅA)", "Summa effekt (ÅA)",
        "Heroma MA 2025-12", "Heroma MA 2026-04", "Heroma MA 2026-08",
        "Inlånade internt PR T3 2025", "Inlånade internt PR T1 2026", "Inlånade internt PR T2 2026",
        "Inlånade TRV T3 2025", "Inlånade TRV T1 2026", "Inlånade TRV T2 2026"]], rows)

# förändring resursbehov T2/T3/T1 (avdelningskod x år x anställd/konsult)
ws = wb["Förändring resursbehov T2,T3,T1"]
rows = []
for top in (7, 22, 37):
    label = re.search(r"T\d\s+\d{4}", s(ws.cell(top - 3, 1).value)).group(0)
    years = {c: s(ws.cell(top - 1, c).value) for c in range(2, 14) if ws.cell(top - 1, c).value}
    for r in range(top + 1, top + 10):
        a = s(ws.cell(r, 1).value)
        if not a or a == "Totalsumma":
            continue
        for ystart, y in years.items():
            for off, typ in ((0, "Anställd"), (1, "Konsult")):
                v = num(ws.cell(r, ystart + off).value)
                if v:
                    rows.append([label, a, y, typ, v])
add("forandring_resursbehov", "Inrapporterat resursbehov per prognostillfälle (T2 2025, T3 2025, T1 2026)", src,
    "Totalt resursbehov (ÅA) per verksamhetsområde, år och Anställd/Konsult – för att jämföra hur behovet ändrats mellan tillfällen.",
    [col("Prognostillfälle", role="dim"), col("Verksamhetsområde", role="dim"), col("År", role="dim"),
     col("Resurstyp", role="dim"), col("ÅA", "num")], rows)

# rekryteringsbehov 2027
ws = wb["Resurs o rekryteringsbehov 2027"]
rows = []
for r in range(6, 16):
    k = s(ws.cell(r, 2).value)
    if k:
        rows.append([k] + [num(ws.cell(r, c).value) or 0 for c in (3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13)])
add("rekrytering_2027", "Resurs- och rekryteringsbehov 2027 per befattningskategori", src,
    "Utfall (Heroma), resursbehov 2026–2028, ökat/minskat behov, personalomsättning och rekryteringsbehov.",
    [col("Befattningskategori", role="dim")] + [col(n, "num") for n in [
        "Utfall 2025", "Utfall jan-apr 2026", "Utfall jan-jul 2026", "Resursbehov 2026", "Resursbehov 2027",
        "Resursbehov 2028", "+/- 2027 vs utfall 2025", "+/- 2027 vs utfall jan-jul 2026",
        "Extern personalomsättning", "Omsättning (ÅA)", "Rekryteringsbehov t.o.m. 2027"]], rows)

# befattning per kategori (mappning)
ws = wb["Befattning per Kategori"]
rows, cur = [], None
for r in range(4, ws.max_row + 1):
    a, b, c = (s(ws.cell(r, i).value) for i in (1, 2, 3))
    if a:
        cur = a
    for v, typ in ((b, "Intern PR"), (c, "Inlånad")):
        if v and cur:
            rows.append([cur, v.lstrip("*"), typ])
add("befattning_kategori", "Befattning per befattningskategori (mappning)", src,
    "Vilka befattningar som hör till vilken befattningskategori.",
    [col("Befattningskategori", role="dim"), col("Befattning", role="dim"), col("Typ", role="dim")], rows)

# ---------------------------------------------------------------- resursanalyser
def resursfil(part):
    f = find(part)
    return f, os.path.basename(f)

# T2: resursinformation per resurs och år
f_t2, n_t2 = resursfil("T2_251006")
wb2 = openpyxl.load_workbook(f_t2, read_only=True, data_only=True)
rows = []
ws = wb2["Resursinformation"]
years = None
for i, r in enumerate(ws.iter_rows(min_row=9, values_only=True)):
    if i == 0:
        years = [(j, str(r[j])) for j in range(11, 19)]
        continue
    if not r[2]:
        continue
    if str(r[2]).strip() == "Totalsumma":
        continue
    for j, y in years:
        v = num(r[j])
        if v:
            rows.append([s(r[2]), s(r[3]), s(r[4]), s(r[5]), s(r[7]), s(r[8]), s(r[10]), num(r[6]), y, v])
add("resurs_t2", "Resursbehov per resurs och år (T2 2026)", n_t2,
    "Varje resurs (namn/resursnr) med avdelning, enhet, befattning, roll och resurstyp samt resursbehov i ÅA per år 2025–2032.",
    [col("Namn", role="dim"), col("Resursnr", role="dim"), col("Resursens avdelning", role="dim"),
     col("Resursens enhet", role="dim"), col("Befattning", role="dim"), col("Roll", role="dim"),
     col("Resurstyp", role="dim"), col("Timkostnad", "num"), col("År", role="dim"), col("ÅA", "num")], rows)

# T2: MA per befattning & roll (månader)
ws = wb2["Resursbehov MA befat. & roll"]
it = list(ws.iter_rows(min_row=9, values_only=True))
months = [(j, str(h)) for j, h in enumerate(it[0]) if j >= 4 and h]
rows = []
for r in it[1:]:
    if not r[2] or str(r[2]).strip() == "Totalsumma":
        continue
    for j, m in months:
        v = num(r[j])
        if v:
            rows.append([s(r[2]), s(r[3]), f"{m[:4]}-{m[4:]}", v])
add("befattning_ma_t2", "Resursbehov per befattning och månad (T2 2026, MA)", n_t2,
    "Månadsarbetskrafter (MA) per befattning och roll, januari 2025 – december 2027.",
    [col("Befattning", role="dim"), col("Roll", role="dim"), col("Månad", role="dim"), col("MA", "num")], rows)

def program_tabell(wbx, sheet, tillfalle, srcname, hdr_row):
    ws = wbx[sheet]
    it = list(ws.iter_rows(min_row=hdr_row, max_row=hdr_row + 40, values_only=True))
    yrs = {}
    for j, v in enumerate(it[0]):
        if v and re.fullmatch(r"\d{4}", str(v)):
            yrs[j] = str(v)
    rows = []
    for r in it[2:]:
        p = s(r[2])
        if not p or p.upper().startswith("TOTAL"):
            continue
        for j, y in yrs.items():
            vak, til = num(r[j]) or 0, num(r[j + 1]) or 0
            if vak:
                rows.append([tillfalle, p, y, "Vakant", vak])
            if til:
                rows.append([tillfalle, p, y, "Tilldelat", til])
    return rows

rows = program_tabell(wb2, "Resursbehov program & befat. ", "T2 2026", n_t2, 7)
# T1 / T3
for part, label in (("T1_2026", "T1 2026"), ("T3_2025", "T3 2025")):
    f, n = resursfil(part)
    w = openpyxl.load_workbook(f, read_only=True, data_only=True)
    hdr = 5 if label == "T1 2026" else 4
    rows += program_tabell(w, "Resursbehov program & befat. ", label, n, hdr)
add("program_aa", "Resursbehov per program och år (T3 2025, T1 2026, T2 2026)", n_t2,
    "ÅA per program och år 2026–2032, uppdelat Vakant/Tilldelat, för tre prognostillfällen.",
    [col("Prognostillfälle", role="dim"), col("Program", role="dim"), col("År", role="dim"),
     col("Status", role="dim"), col("ÅA", "num")], rows)

# T1 & T3: egna resurser per år (kategori), flerradig header
def egna(part, label):
    f, n = resursfil(part)
    w = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = w["Egna resurser År"]
    it = list(ws.iter_rows(min_row=1, max_row=40, values_only=True))
    hdr_i = next(i for i, r in enumerate(it) if r[1] and str(r[1]).startswith("Befattningskategori"))
    grp_i = hdr_i - 1
    groups, g = {}, None
    for j in range(2, 14):
        if it[grp_i][j]:
            g = str(it[grp_i][j]).strip()
        groups[j] = g
    out = []
    for r in it[hdr_i + 1:]:
        k = s(r[1])
        if not k or k == "Totalsumma":
            if k == "Totalsumma":
                break
            continue
        for j in range(2, 14):
            v = num(r[j])
            if v:
                out.append([label, k, groups[j], str(it[hdr_i][j]), v])
    return out, n

rows = []
for part, label in (("T3_2025", "T3 2025"), ("T1_2026", "T1 2026")):
    o, n = egna(part, label)
    rows += o
add("egna_resurser", "Egna resurser: nuvarande bemanning och planerade rekryteringar (T3 2025, T1 2026)", n,
    "ÅA per befattningskategori och år: nuvarande bemanning (anställda/konsulter) samt planerade rekryteringar.",
    [col("Prognostillfälle", role="dim"), col("Befattningskategori", role="dim"), col("Mått", role="dim"),
     col("År", role="dim"), col("ÅA", "num")], rows)

# ---------------------------------------------------------------- PowerPoint
from pptx import Presentation
f_pp = find("Presentation_av_uppdrag")
prs = Presentation(f_pp)
rows = []
for idx, sl in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text.strip() for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    texts = [t for t in texts if not t.startswith("Trafikverkets interna")]
    if not texts:
        continue
    avd = texts[-1] if len(texts[-1]) < 40 else None
    body = texts[:-1] if avd else texts
    typ = "Övrigt"
    flat = " ".join(texts)
    if "Risker/Möjligheter" in flat:
        typ = "Generell bedömning – risker/möjligheter"
    elif "Bubblare" in flat:
        typ = "Bubblare"
    elif "Summerad effekt" in flat:
        typ = "Summerad effekt (bild)"
    elif "Föreslagna åtgärder" in flat:
        typ = "Föreslagna åtgärder (bild)"
    text = "\n".join(t for t in body if t not in ("Generell bedömning!", "Risker/Möjligheter", "Bubblare!"))
    if text.strip():
        rows.append([idx, avd, typ, text])
add("slides", "Presentationens text per avdelning (slides)", os.path.basename(f_pp),
    "Löpande text, risker/möjligheter och bubblare per avdelning från PowerPoint-sammanställningen.",
    [col("Slide", "num"), col("Avdelning", role="dim"), col("Typ", role="dim"), col("Text", role="text")], rows)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("window.PRELOADED = ")
    json.dump(DATASETS, fh, ensure_ascii=False, separators=(",", ":"))
    fh.write(";\n")
print("Skrev", OUT, round(os.path.getsize(OUT) / 1024), "kB")
